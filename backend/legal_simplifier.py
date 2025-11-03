import os
from groq import Groq
from dotenv import load_dotenv
import PyPDF2
import docx
import json
from datetime import datetime
from fpdf import FPDF

load_dotenv('backend/.env')

class LegalDocumentSimplifier:
    def __init__(self):
        """
        Initialize the Legal Document Simplifier
        Uses Groq's free API with Llama models
        """
        self.api_key = os.environ.get("GROQ_API_KEY")
        if not self.api_key:
            print("\n⚠️  GROQ_API_KEY not found!")
            print("Get your free API key at: https://console.groq.com")
            print("Then set it: export GROQ_API_KEY='your-key-here'")
            raise ValueError("GROQ_API_KEY environment variable not set")
        
        self.client = Groq(api_key=self.api_key)
        self.conversation_history = []
        self.loaded_documents = {}  
        self.current_analysis = None
        
    def read_pdf(self, file_path):
        """Extract text from PDF file"""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text()
                return text
        except Exception as e:
            return f"Error reading PDF: {str(e)}"
    
    def read_docx(self, file_path):
        """Extract text from DOCX file"""
        try:
            doc = docx.Document(file_path)
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            return text
        except Exception as e:
            return f"Error reading DOCX: {str(e)}"
    
    def read_txt(self, file_path):
        """Extract text from TXT file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except Exception as e:
            return f"Error reading TXT: {str(e)}"
    
    def load_document(self, file_path, doc_name=None):
        """Load document based on file extension"""
        extension = file_path.lower().split('.')[-1]
        
        if extension == 'pdf':
            text = self.read_pdf(file_path)
        elif extension == 'docx':
            text = self.read_docx(file_path)
        elif extension == 'txt':
            text = self.read_txt(file_path)
        else:
            return "Unsupported file format. Please use PDF, DOCX, or TXT files."
        
        if not doc_name:
            doc_name = os.path.basename(file_path)
        
        self.loaded_documents[doc_name] = {
            'text': text,
            'path': file_path,
            'loaded_at': datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
        
        return text
    
    def highlight_concerning_clauses(self, legal_text):
        """
        Use AI to identify and highlight concerning clauses
        """
        prompt = f"""You are a legal expert analyzing documents for potential concerns. 

Analyze this legal document and identify concerning clauses or red flags. For each concerning clause:
1. Quote the clause (keep it brief)
2. Explain why it's concerning in simple terms
3. Rate the severity (LOW, MEDIUM, HIGH)

Return ONLY a JSON array with this exact structure:
[
  {{
    "clause": "brief quote of concerning text",
    "concern": "explanation in plain language",
    "severity": "HIGH/MEDIUM/LOW",
    "recommendation": "what the reader should do"
  }}
]

Legal Text:
{legal_text[:4000]}

Return ONLY valid JSON, no other text."""

        try:
            response = self.client.chat.completions.create(
                messages=[
                    {
                        "role": "system",
                        "content": "You are a legal expert who identifies concerning clauses. Return only valid JSON."
                    },
                    {
                        "role": "user",
                        "content": prompt
                    }
                ],
                model="llama-3.3-70b-versatile",
                temperature=0.5,
                max_tokens=2000
            )
            
            result = response.choices[0].message.content.strip()
            if "```json" in result:
                result = result.split("```json")[1].split("```")[0].strip()
            elif "```" in result:
                result = result.split("```")[1].split("```")[0].strip()
            
            return json.loads(result)
        
        except Exception as e:
            print(f"Error analyzing concerns: {str(e)}")
            return []
    
    def simplify_legal_text(self, legal_text, query=None):
        """
        Simplify legal text using Groq's API
        """
        if query:
            prompt = f"""You are a legal document simplifier. A user has provided this legal text and has a specific question about it.

Legal Text:
{legal_text[:3000]}

User Question: {query}

Please answer their question in simple, plain language that anyone can understand. Avoid legal jargon and explain any complex terms you must use."""
        else:
            prompt = f"""You are a legal document simplifier. Your job is to break down complex legal language into simple, easy-to-understand terms.

Please analyze this legal text and provide:
1. A brief summary (2-3 sentences)
2. Key points in plain language (numbered list)
3. Important dates or deadlines mentioned
4. Your obligations and rights

Legal Text:
{legal_text[:3000]}

Remember to use everyday language that anyone can understand."""

        try:
            response = self.client.chat.completions.create(
                messages=[
                    {
                        "role": "system",
                        "content": "You are a helpful legal document simplifier. You break down complex legal jargon into plain, simple language that anyone can understand."
                    },
                    {
                        "role": "user",
                        "content": prompt
                    }
                ],
                model="llama-3.3-70b-versatile",
                temperature=0.7,
                max_tokens=1500
            )
            
            return response.choices[0].message.content
        
        except Exception as e:
            return f"Error calling API: {str(e)}"
    
    def compare_documents(self, doc1_name, doc2_name):
        """
        Compare two loaded documents and highlight differences
        """
        if doc1_name not in self.loaded_documents or doc2_name not in self.loaded_documents:
            return "One or both documents not found. Please load them first."
        
        doc1_text = self.loaded_documents[doc1_name]['text'][:3000]
        doc2_text = self.loaded_documents[doc2_name]['text'][:3000]
        
        prompt = f"""Compare these two legal documents and provide:
1. Main differences between them
2. Which document is more favorable and why
3. Key clauses that differ
4. Recommendations on which is better

Document 1 ({doc1_name}):
{doc1_text}

Document 2 ({doc2_name}):
{doc2_text}

Provide a clear, easy-to-understand comparison."""

        try:
            response = self.client.chat.completions.create(
                messages=[
                    {
                        "role": "system",
                        "content": "You are a legal expert who compares documents and explains differences clearly."
                    },
                    {
                        "role": "user",
                        "content": prompt
                    }
                ],
                model="llama-3.3-70b-versatile",
                temperature=0.7,
                max_tokens=2000
            )
            
            return response.choices[0].message.content
        
        except Exception as e:
            return f"Error comparing documents: {str(e)}"
    
    def save_to_pdf(self, content, filename=None):
        """
        Save simplified analysis to PDF
        """
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"simplified_legal_doc_{timestamp}.pdf"
        
        os.makedirs('outputs', exist_ok=True)
        
        if not filename.startswith('outputs/'):
            filename = os.path.join('outputs', filename)
        
        try:
            pdf = FPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=15)
            
            pdf.set_font("Arial", "B", 16)
            pdf.cell(0, 10, "Legal Document Analysis", ln=True, align="C")
            pdf.ln(5)
            
            pdf.set_font("Arial", "I", 10)
            pdf.cell(0, 10, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
            pdf.ln(5)
            
            pdf.set_font("Arial", "", 11)
            
            content = content.encode('latin-1', 'replace').decode('latin-1')
            
            lines = content.split('\n')
            for line in lines:
                if line.strip():
                    pdf.multi_cell(0, 6, line)
                else:
                    pdf.ln(3)
            
            pdf.output(filename)
            return filename
        
        except Exception as e:
            return f"Error saving PDF: {str(e)}"
    
    def chat(self, user_message, document_context=None):
        """
        Handle conversational queries about the document
        """
        if document_context:
            context_prompt = f"Document Context:\n{document_context[:2000]}\n\nUser Question: {user_message}"
        else:
            context_prompt = user_message
        
        self.conversation_history.append({"role": "user", "content": context_prompt})
        
        try:
            response = self.client.chat.completions.create(
                messages=[
                    {
                        "role": "system",
                        "content": "You are a legal document expert who explains things in simple terms. Be conversational, friendly, and avoid legal jargon."
                    }
                ] + self.conversation_history,
                model="llama-3.3-70b-versatile",
                temperature=0.7,
                max_tokens=1000
            )
            
            assistant_message = response.choices[0].message.content
            self.conversation_history.append({"role": "assistant", "content": assistant_message})
            
            return assistant_message
        
        except Exception as e:
            return f"Error: {str(e)}"